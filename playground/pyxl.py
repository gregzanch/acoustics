from datetime import date

from openpyxl import Workbook
from openpyxl.chart import (
    LineChart,
    Reference,
)
from openpyxl.chart.axis import DateAxis

wb = Workbook()
ws = wb.active

rows = [
  [ 'Frequency', 'Pink Noise', 'A.txt', 'B.txt', 'C.txt' ],
  [ '25', 45.5, 40.3, 40.8, 41.8 ],
  [ '32', 50.3, 49.1, 47.6, 44.4 ],
  [ '40', 49.9, 49.1, 50.1, 47.4 ],
  [ '50', 44.2, 42.4, 43.7, 42.8 ],
  [ '63', 51.5, 51.1, 50.3, 49.5 ],
  [ '80', 58.7, 57.8, 56.2, 55.8 ],
  [ '100', 61.8, 61.1, 58.7, 58.2 ],
  [ '125', 61.8, 61.5, 58.3, 57.8 ],
  [ '160', 64.8, 64.7, 62.2, 61.6 ],
  [ '200', 66.9, 67.1, 65.3, 64.5 ],
  [ '250', 67.7, 68, 65.8, 65.3 ],
  [ '315', 67.8, 67.3, 65, 64.4 ],
  [ '400', 67.9, 67.4, 64.2, 64.6 ],
  [ '500', 65.9, 65.8, 62.1, 62.6 ],
  [ '630', 65.2, 64.9, 63.7, 63.9 ],
  [ '800', 67.2, 67.4, 66.6, 67.2 ],
  [ '1000', 71.6, 71.2, 70.5, 70.6 ],
  [ '1250', 71.6, 71.6, 70.3, 70.6 ],
  [ '1600', 71.7, 71.9, 70.7, 70.8 ],
  [ '2000', 75.2, 75.1, 74, 74.2 ],
  [ '2500', 75.7, 75.8, 73.8, 74 ],
  [ '3150', 76, 76, 72.6, 72.5 ],
  [ '4000', 75.9, 75.3, 70.8, 70.6 ],
  [ '5000', 74.4, 73.1, 67.5, 67 ],
  [ '6300', 73.4, 70.1, 62.7, 63.1 ],
  [ '8000', 71.4, 67.5, 63.4, 63.2 ],
  [ '10000', 69.1, 66.7, 66.5, 66.5 ],
  [ '12500', 71.4, 71.7, 68.2, 67.9 ],
  [ '16000', 72.3, 69.9, 65.9, 66.4 ],
  [ '20000', 71.1, 68.2, 64.1, 64.1 ]
]
for row in rows:
    ws.append(row)

c1 = LineChart()
c1.title = "Line Chart"
c1.style = 13
c1.y_axis.title = 'dB'
c1.x_axis.title = 'Frequency (Hz)'

data = Reference(ws, min_col=2, min_row=1, max_col=4, max_row=7)
c1.add_data(data, titles_from_data=True)

# Style the lines
s1 = c1.series[0]
s1.marker.symbol = "triangle"
s1.marker.graphicalProperties.solidFill = "FF0000" # Marker filling
s1.marker.graphicalProperties.line.solidFill = "FF0000" # Marker outline

s1.graphicalProperties.line.noFill = True

s2 = c1.series[1]
s2.graphicalProperties.line.solidFill = "00AAAA"
s2.graphicalProperties.line.dashStyle = "sysDot"
s2.graphicalProperties.line.width = 100050 # width in EMUs

s2 = c1.series[2]
s2.smooth = True # Make the line smooth

ws.add_chart(c1, "A10")

from copy import deepcopy
stacked = deepcopy(c1)
stacked.grouping = "stacked"
stacked.title = "Stacked Line Chart"
ws.add_chart(stacked, "A27")

percent_stacked = deepcopy(c1)
percent_stacked.grouping = "percentStacked"
percent_stacked.title = "Percent Stacked Line Chart"
ws.add_chart(percent_stacked, "A44")

# Chart with date axis
c2 = LineChart()
c2.title = "Date Axis"
c2.style = 12
c2.y_axis.title = "Size"
c2.y_axis.crossAx = 500
c2.x_axis = DateAxis(crossAx=100)
c2.x_axis.number_format = 'd-mmm'
c2.x_axis.majorTimeUnit = "days"
c2.x_axis.title = "Date"

c2.add_data(data, titles_from_data=True)
dates = Reference(ws, min_col=1, min_row=2, max_row=7)
c2.set_categories(dates)

ws.add_chart(c2, "A61")

wb.save("line.xlsx")
